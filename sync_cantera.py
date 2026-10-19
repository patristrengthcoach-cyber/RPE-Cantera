#!/usr/bin/env python3
"""
Sincroniza el dashboard VALORACIÓN CANTERA con los Excel de Drive.

Lee el/los "Registro Valoraciones" (una hoja por categoría) y el roster ID_JUGADOR_RCF,
detecta mediciones NUEVAS (fecha posterior a la última que ya conoce el dashboard) y
actualiza los datos embebidos en dashboard.html:

  · altura (y alturaPrev / histórico de alturas)         → todas las categorías
  · medición antropométrica completa (perímetros, ROM)   → asimetrías, avisos de riesgo
  · velocidad de crecimiento, alertas PHV
  · percentiles por categoría y curva de crecimiento
  · jugadores nuevos (nombre/equipo desde el roster)
  · nombres / equipo del roster
  · pliegues cutáneos (Σ8, % graso) y la altura que traigan esos archivos

Es idempotente: si no hay nada nuevo no toca el archivo.
Los archivos de "Pliegues" también se leen: se vinculan a cada jugador por ID o, si el Excel solo trae el nombre,
automáticamente por nombre con el roster (nunca se une un nombre repetido bajo dos IDs).

Uso local:
    python sync_cantera.py --html dashboard.html --registro registro.xlsx [--registro otro.xlsx] --roster roster.xlsx
Uso con Drive (cuenta de servicio):
    python sync_cantera.py --html dashboard.html --config sync_config.json --sa-json cuenta.json
"""
import argparse
import copy
import datetime as dt
import difflib
import io
import json
import os
import re
import sys
import unicodedata
from collections import Counter

import numpy as np
import openpyxl
from scipy.stats import percentileofscore

DATASET_RE = re.compile(r'(<script id="dataset" type="application/json">)(.*?)(</script>)', re.S)

RAW_KEYS = [
    'Cuello',
    'Brazo DERECHO relajado (1/2 brazo)', 'Brazo IZQUIERDO relajado (1/2 brazo)',
    'Brazo DERECHO 90º (MAX)', 'Brazo IZQUIERDO 90º (MAX)',
    'Antebrazo DERECHO (MAX)', 'Antebrazo IZQUIERDO (MAX)',
    'Pecho / Tórax (bajo axilas)', 'Perimetro Cadera/Gluteo',
    'Muslo DERECHO (1cm bajo gluteo MAX)', 'Muslo IZQUIERDO (1cm bajo gluteo MAX)',
    'Muslo Medio DERECHO (split cajón/ingle-rotula 1/2)', 'Muslo Medio IZQUIERDO (split cajón/ingle-rotula 1/2)',
    'Pierna DERECHA (MAX)', 'Pierna IZQUIERDA (MAX)',
    'Tobillo DERECHO (MIN)', 'Tobillo IZQUIERDO (MIN)',
    'Dorsiflexion Tobillo DERECHO', 'Dorsiflexion Tobillo IZQUIERDO',
    'Rotación INTERNA Cadera DERECHA (peroneo)', 'Rotación INTERNA Cadera IZQUIERDA',
    'Rotación EXTERNA Cadera DERECHA', 'Rotación EXTERNA Cadera IZQUIERDA',
]
ASYM_PAIRS = {
    'brazo_relajado': (1, 2), 'brazo_90': (3, 4), 'antebrazo': (5, 6), 'muslo': (9, 10),
    'muslo_medio': (11, 12), 'pierna': (13, 14), 'tobillo': (15, 16), 'dorsiflexion': (17, 18),
    'rot_interna_cadera': (19, 20), 'rot_externa_cadera': (21, 22),
}
ASYM_NAMES = {
    'brazo_relajado': 'Brazo relajado', 'brazo_90': 'Brazo (90°)', 'antebrazo': 'Antebrazo', 'muslo': 'Muslo',
    'muslo_medio': 'Muslo medio', 'pierna': 'Pierna', 'tobillo': 'Tobillo', 'dorsiflexion': 'Dorsiflexión tobillo',
    'rot_interna_cadera': 'Rotación interna cadera', 'rot_externa_cadera': 'Rotación externa cadera',
}
SHEET_CATEGORY = {
    'JUV_A': 'Juvenil A', 'JUV_B': 'Juvenil B', 'CAD_A': 'Cadete A', 'CAD_B': 'Cadete B',
    'INF_A': 'Infantil A', 'INF_B': 'Infantil B', 'INF_FEM': 'Infantil Femenino',
    'SENIOR FEM': 'Senior Femenino', 'ALEVIN': 'Alevín',
}
PHV_TXT = {
    'pre': ('Antes de PHV', 'Crecimiento rápido antes de PHV — vigilar patología de Sever (talón).', 'medium'),
    'near': ('Cerca de / en PHV', 'Cerca o en PHV — vigilar patología de rodilla (ej. Osgood-Schlatter).', 'medium'),
    'post': ('Después de PHV', 'Después de PHV con crecimiento sostenido por encima del umbral — patrón atípico, tratar como riesgo.', 'high'),
}
TRAIN_NOTE = 'Técnica motriz y reducción de volumen de COD / ACC / DECC.'


# ----------------------------------------------------------------------------- utilidades
def norm(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9 ]', '', s).strip()


def dmy(s):
    return dt.datetime.strptime(s, '%d/%m/%Y').date()


def fmt_date(d):
    return d.strftime('%d/%m/%Y')


def parse_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for f in ('%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d', '%d-%m-%Y', '%d.%m.%Y'):
            try:
                return dt.datetime.strptime(s, f).date()
            except ValueError:
                pass
    return None


def parse_num(v):
    """Número de una celda. Google Sheets en configuración regional distinta convierte '23,5' en la
    fecha 23-may: se deshace (día + mes/10; mes>=10 -> mes/100)."""
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return round(v.day + (v.month / 10 if v.month < 10 else v.month / 100), 2)
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '.')
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def clean(o):
    if isinstance(o, float) and (o != o or o in (float('inf'), float('-inf'))):
        return None
    if isinstance(o, dict):
        return {k: clean(v) for k, v in o.items()}
    if isinstance(o, list):
        return [clean(v) for v in o]
    return o


def pct_rank(vals, v):
    return round(float(percentileofscore(vals, v, kind='mean')), 1)


# ----------------------------------------------------------------------------- lectura de Excel
def read_registro(xlsx_bytes_or_path, log):
    """Devuelve lista de filas: dict(sheet,id,fecha(date),categoria,altura,peso,vals{rawkey:val}).
    Las columnas se identifican por el texto de la cabecera (cada archivo puede traer un orden distinto)."""
    wb = openpyxl.load_workbook(xlsx_bytes_or_path, data_only=True)
    raw_by_norm = {norm(k): k for k in RAW_KEYS}
    rows = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration:
            continue
        col = {}                                   # nombre lógico -> índice de columna
        raw_cols = {}                              # índice -> clave raw
        for i, h in enumerate(header):
            n = norm(h or '')
            if not n:
                continue
            if n.startswith('fecha') and 'fecha' not in col:
                col['fecha'] = i
            elif n.startswith('categoria') and 'cat' not in col:
                col['cat'] = i
            elif n.startswith('id de') and 'id' not in col:
                col['id'] = i
            elif n.startswith('altura') and 'altura' not in col:
                col['altura'] = i
            elif n.startswith('peso') and 'peso' not in col:
                col['peso'] = i
            elif n in raw_by_norm:
                raw_cols[i] = raw_by_norm[n]
        if 'fecha' not in col or 'id' not in col or 'altura' not in col:
            log(f'  hoja "{ws.title}" ignorada (no tiene el formato de Registro Valoraciones)')
            continue
        cat_default = SHEET_CATEGORY.get(ws.title.strip().upper().replace('Í', 'I'))
        cur_date, cur_cat = None, cat_default
        get = lambda r, k: r[col[k]] if k in col and col[k] < len(r) else None
        for r in it:
            r = list(r)
            d = parse_date(get(r, 'fecha'))
            if d:
                cur_date = d
            c = get(r, 'cat')
            if isinstance(c, str) and c.strip():
                cur_cat = c.strip()
            pid = get(r, 'id')
            pid = str(pid).strip() if pid is not None else ''
            if not pid or cur_date is None:
                continue
            vals = {}
            for i, k in raw_cols.items():
                v = parse_num(r[i]) if i < len(r) else None
                if v is not None:
                    vals[k] = v
            rows.append(dict(sheet=ws.title, id=pid, fecha=cur_date, categoria=cur_cat,
                             altura=parse_num(get(r, 'altura')), peso=parse_num(get(r, 'peso')), vals=vals))
    return rows


def read_roster(xlsx_bytes_or_path):
    wb = openpyxl.load_workbook(xlsx_bytes_or_path, data_only=True)
    ws = wb['ID_JUGADOR'] if 'ID_JUGADOR' in wb.sheetnames else wb.worksheets[0]
    out = {}
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not r or not r[0]:
            continue
        pid = str(r[0]).strip()
        name = ' '.join(f'{r[1] or ""} {r[2] or ""}'.split())
        team = str(r[3]).strip() if len(r) > 3 and r[3] else None
        if name:
            out[pid] = dict(name=name, team=team)
    return out


# ----------------------------------------------------------------------------- pliegues (lectura)
FOLDS = ['triceps', 'biceps', 'subescapular', 'suprailiaca', 'supraespinal', 'abdominal', 'musloAnterior', 'piernaMedial']
FOLD_SYN = {
    'triceps': ('triceps', 'tricipital'),
    'biceps': ('biceps', 'bicipital'),
    'subescapular': ('subescapular', 'subescap'),
    'suprailiaca': ('suprailiaca', 'suprailiaco', 'suprailiac', 'cresta iliaca'),
    'supraespinal': ('supraespinal', 'supraespinoso', 'oblicuo'),
    'abdominal': ('abdominal', 'abdomen'),
    'musloAnterior': ('muslo anterior', 'muslo'),
    'piernaMedial': ('pierna medial', 'pierna', 'gemelo'),
}
DATE_IN_TEXT = re.compile(r'(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})')
RCF_RE = re.compile(r'RCF[\s_-]*(\d+)', re.I)
JUNK_NAMES = {'total', 'totales', 'media', 'promedio', 'desv', 'sumatorio', 'maximo', 'minimo', 'max', 'min'}


def _date_from_text(s):
    m = DATE_IN_TEXT.search(str(s))
    if not m:
        return None, str(s)
    d, mo, y = (int(x) for x in m.groups())
    if y < 100:
        y += 2000
    try:
        return dt.date(y, mo, d), (str(s)[:m.start()] + ' ' + str(s)[m.end():])
    except ValueError:
        return None, str(s)


def _classify_header(h):
    """-> (clave, fecha_en_la_cabecera). Claves: nombre, id, fecha, cat, altura, sum, o uno de FOLDS."""
    if h is None:
        return None, None
    if isinstance(h, (dt.datetime, dt.date)):
        return 'nombre', parse_date(h)
    hd, txt = _date_from_text(h)
    n = norm(txt)
    if not n:
        return ('nombre' if hd else None), hd
    if n.startswith('nombre') or n in ('jugador', 'jugadora', 'name', 'apellidos y nombre'):
        return 'nombre', hd
    if n.startswith('fecha'):
        return 'fecha', hd
    if n.startswith('categoria'):
        return 'cat', hd
    if n == 'id' or n.startswith('id ') or n.startswith('idj'):
        return 'id', hd
    if n.startswith('altura') or n.startswith('talla'):
        return 'altura', hd
    if n.startswith('sumatorio') or n in ('suma', 'total', 'suma pliegues', 'sum'):
        return 'sum', hd
    for k, syns in FOLD_SYN.items():
        if any(n == s or n.startswith(s + ' ') for s in syns):
            return k, hd
    return None, hd


def read_pliegues(xlsx_bytes_or_path, hint=None, log=print):
    """Lee un archivo de pliegues (cualquier hoja con cabecera de pliegues). Admite:
      · jugadores por ID (RCF_xxx) o por NOMBRE (con o sin columna ID), con o sin columna Categoría
      · la fecha en una columna «Fecha», en la celda de cabecera (p. ej. «28-09-2026 Nombre») o en filas superiores
      · columna «Altura» opcional
    Devuelve filas: dict(sheet,id,name,fecha,categoria,altura,folds,sumFile,hint)."""
    wb = openpyxl.load_workbook(xlsx_bytes_or_path, data_only=True)
    out = []
    for ws in wb.worksheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        hi, col = None, {}
        for i, row in enumerate(grid[:25]):
            cl = {}
            for j, h in enumerate(row):
                k, _ = _classify_header(h)
                if k and k not in cl:
                    cl[k] = j
            if sum(1 for f in FOLDS if f in cl) >= 4:
                hi, col = i, cl
                break
        if hi is None:
            log(f'  hoja "{ws.title}" ignorada (no tiene cabecera de pliegues)')
            continue
        if 'id' not in col and 'nombre' not in col:
            log(f'  hoja "{ws.title}" ignorada (sin columna de ID ni de nombre)')
            continue
        # fecha por defecto: cabecera -> filas superiores
        default_date = None
        for h in grid[hi]:
            hd = _classify_header(h)[1]
            if hd:
                default_date = hd
                break
        if default_date is None:
            for row in grid[:hi]:
                for c in row:
                    d = parse_date(c) or (_date_from_text(c)[0] if isinstance(c, str) else None)
                    if d:
                        default_date = d
        cur_date, cur_cat = None, None
        for row in grid[hi + 1:]:
            get = lambda k: row[col[k]] if k in col and col[k] < len(row) else None
            d = parse_date(get('fecha'))
            if d:
                cur_date = d
            c = get('cat')
            if isinstance(c, str) and c.strip():
                cur_cat = c.strip()
            raw_id, raw_name = '', ''
            for k in ('id', 'nombre'):
                v = get(k)
                if v is None or isinstance(v, (int, float, dt.datetime, dt.date)):
                    continue
                s = ' '.join(str(v).split())
                if not s:
                    continue
                m = RCF_RE.fullmatch(s)
                if m:
                    raw_id = raw_id or f'RCF_{m.group(1).zfill(3)}'
                elif not raw_name:
                    raw_name = s
            if not raw_id and not raw_name:
                continue
            if not raw_id and norm(raw_name) in JUNK_NAMES:
                continue
            folds = {f: (parse_num(get(f)) if f in col else None) for f in FOLDS}
            if all(v is None for v in folds.values()):
                continue
            fecha = cur_date or default_date
            if fecha is None:
                log(f'  fila «{raw_id or raw_name}» de «{ws.title}» sin fecha — ignorada')
                continue
            out.append(dict(sheet=ws.title, id=raw_id, name=raw_name, fecha=fecha, categoria=cur_cat,
                            altura=parse_num(get('altura')), folds=folds, sumFile=parse_num(get('sum')), hint=hint))
    return out


# ----------------------------------------------------------------------------- vinculación nombre -> ID
class Resolver:
    """Asocia nombres a IDs RCF usando el roster y los nombres que ya tiene el dashboard.
    Un nombre que el roster repite bajo ≥2 IDs NUNCA se vincula por nombre (salvo que la categoría
    del archivo desempate): se mantienen como personas distintas."""

    def __init__(self, data, roster):
        self.P = {p['id']: p for p in data['players']}
        self.names, self.cats, self.eff = {}, {}, {}
        for p in data['players']:
            self._add(p['id'], p['name'], [p['anthro']['categoria'], p.get('equipoOficial')])
            if not p.get('isUnresolvedId'):               # los IDs provisionales (Alevín) no cuentan como duplicados
                self.eff[p['id']] = p['name']
        for pid, r in (roster or {}).items():
            self._add(pid, r['name'], [r['team']])
            self.eff[pid] = r['name']                      # el roster manda
        self.by_key, self.by_flat = {}, {}
        for pid, ns in self.names.items():
            for n in ns:
                t = norm(n).split()
                self.by_key.setdefault(tuple(sorted(t)), set()).add(pid)
                self.by_flat.setdefault(''.join(t), set()).add(pid)
        g = {}
        for pid, n in self.eff.items():
            g.setdefault(tuple(sorted(norm(n).split())), []).append(pid)
        self.dup_groups = {k: sorted(v) for k, v in g.items() if len(v) > 1 and k}

    def _add(self, pid, name, cats):
        self.names.setdefault(pid, set()).add(name)
        self.cats.setdefault(pid, set()).update(norm(c) for c in cats if c)

    def is_dup_name(self, name):
        return tuple(sorted(norm(name).split())) in self.dup_groups

    def display(self, pid):
        return self.eff.get(pid)

    def cat_ok(self, pid, hint):
        hn = norm(hint)
        return any(c.startswith(hn) for c in self.cats.get(pid, ()))

    def _pick(self, ids, hint):
        ids = set(ids)
        if len(ids) > 1 and hint:
            f = {i for i in ids if self.cat_ok(i, hint)}
            if f:
                ids = f
        return ids

    def lookup(self, raw, hint=None):
        """-> (id|None, how, candidatos). how: 'roster' (nombre exacto) | 'name' (aproximado) | 'ambiguo' | 'sin_coincidencia'."""
        toks = norm(raw).split()
        if not toks:
            return None, 'sin_coincidencia', []
        key = tuple(sorted(toks))
        stages = [('roster', self.by_key.get(key, set())), ('roster', self.by_flat.get(''.join(toks), set()))]
        if len(toks) >= 2:                                 # escribió menos apellidos que el roster
            stages.append(('name', {i for k, ids in self.by_key.items() if set(toks) <= set(k) for i in ids}))
        for how, ids in stages:
            if not ids:
                continue
            ids = self._pick(ids, hint)
            if len(ids) == 1:
                return next(iter(ids)), how, []
            return None, 'ambiguo', sorted(ids)
        keys = [' '.join(k) for k in self.by_key]
        close = difflib.get_close_matches(' '.join(key), keys, n=2, cutoff=0.88)
        if len(close) == 1:
            ids = self._pick({i for k, v in self.by_key.items() if ' '.join(k) == close[0] for i in v}, hint)
            if len(ids) == 1:
                return next(iter(ids)), 'name', []
        return None, 'sin_coincidencia', []

    def resolve_id(self, raw, hint=None):
        if RCF_RE.fullmatch(raw.strip()) or raw in self.P:
            return raw, 'id', []
        return self.lookup(raw, hint)


# ----------------------------------------------------------------------------- cálculos
def asym(a, b):
    if a is None or b is None:
        return None
    m = (a + b) / 2
    return None if m == 0 else round((a - b) / m * 100, 1)


def derive_anthro(a):
    """Recalcula campos derivados de a['raw'] (promedios, asimetrías)."""
    raw = a['raw']
    g = lambda i: raw.get(RAW_KEYS[i])
    a['pecho'], a['cadera'], a['cuello'] = g(7), g(8), g(0)
    mm = [g(11), g(12)]
    a['muslo_medio_avg'] = None if None in mm else round(sum(mm) / 2, 1)
    df = [g(17), g(18)]
    a['dorsiflex_avg'] = None if None in df else round(sum(df) / 2, 1)
    a['asimetrias'] = {k: asym(g(i), g(j)) for k, (i, j) in ASYM_PAIRS.items()}


def antro_flags(a):
    out = []
    for k, v in a['asimetrias'].items():
        if v is None or abs(v) < 15:
            continue
        sev = 'high' if abs(v) >= 25 else 'medium'
        out.append(dict(source='antro', metric=ASYM_NAMES[k], value=v,
                        label=f"{ASYM_NAMES[k]}: {abs(v)}% mayor en lado {'D' if v > 0 else 'I'}", severity=sev))
    return out


def latest_known(p):
    a = p['anthro']
    ds = [parse_date(a.get('fecha')), parse_date(a.get('alturaFecha'))]
    ds = [d for d in ds if d]
    return max(ds) if ds else None


def recompute_growth(p, thr):
    """Velocidad de crecimiento. Con línea base histórica (10/08/2025) se usa esa; sin ella, se ancla en la
    primera altura registrada (growth.derived = True). Los intervalos <90 días se marcan como cortos y no
    disparan alerta PHV (anualizar 2-3 meses amplifica el error de medición)."""
    g = p.get('growth')
    if not g:
        hist = sorted(p.get('alturas') or [], key=lambda x: dmy(x['fecha']))
        if len(hist) < 2:
            return
        g = dict(prevHeight=hist[0]['valor'], prevDate=hist[0]['fecha'], derived=True)
    if p['anthro']['altura'] is None:
        return
    cur = p['anthro'].get('alturaFecha') or p['anthro']['fecha']
    days = (dmy(cur) - dmy(g['prevDate'])).days
    if days <= 0:
        return
    delta = round(p['anthro']['altura'] - g['prevHeight'], 1)
    vel = round(delta / days * 365.25, 1)
    short = bool(g.get('derived')) and days < 90
    ng = dict(prevHeight=g['prevHeight'], prevDate=g['prevDate'], currHeight=p['anthro']['altura'],
              currDate=cur, deltaCm=delta, daysBetween=days, velocityCmYear=vel)
    if g.get('derived'):
        ng['derived'] = True
        if short:
            ng['short'] = True
    p['growth'] = ng
    p['riskFlags'] = [f for f in p['riskFlags'] if f.get('source') != 'phv']
    p['phvAlert'] = None
    st = p.get('phvStage')
    if vel >= thr and st in PHV_TXT and not short:
        lab, note, sev = PHV_TXT[st]
        p['phvAlert'] = dict(stage=st, stageLabel=lab, velocityCmYear=vel, pathologyNote=note, trainingNote=TRAIN_NOTE, severity=sev)
        p['riskFlags'].append(dict(source='phv', metric='Velocidad de crecimiento (PHV)', value=vel,
                                   label=f'{vel} cm/año (>0,6 cm/mes) · {lab} · {note}', severity=sev))
    p['riskCount'] = len(p['riskFlags'])


# ----------------------------------------------------------------------------- sincronización
def _set_auto(lst, kind, items):
    """Sustituye en `lst` los avisos generados automáticamente (campo auto=kind). True si cambió el contenido."""
    old = {x['text'] for x in lst if x.get('auto') == kind}
    new = {t for _, t in items}
    if old == new:
        return False
    lst[:] = [x for x in lst if x.get('auto') != kind] + [dict(level=l, text=t, auto=kind) for l, t in items]
    return True


def _pl_record_from_player(p):
    q, a = p['pliegues'], p['anthro']
    return dict(playerId=p['id'], id=p['id'], name=p['name'], categoria=a['categoria'], fecha=q['fecha'],
                altura=a['altura'], alturaFecha=a.get('alturaFecha') or a['fecha'],
                alturaPrev=(a.get('alturaPrev') or {}).get('valor'),
                **{k: q[k] for k in FOLDS}, sum8=q['sumatorio'], pctFat=q['pctFatEst'], matchedBy=q['matchedBy'])


def sync(data, registro_list, roster, log=print, pliegues=None):
    """Modifica `data` in situ. Devuelve lista de cambios (texto).
    registro_list: filas de read_registro; pliegues: filas de read_pliegues."""
    T = data['team']
    P = {p['id']: p for p in data['players']}
    thr = T['phvThresholdCmYear']
    changes, touched_cats, issues = [], set(), []
    now_iso = lambda: dt.datetime.now().strftime('%d/%m/%Y %H:%M')
    pl = data.setdefault('pliegues', dict(records=[], notes=[], equation='Slaughter 1988 · tríceps + pierna medial · masculina'))
    pl.setdefault('records', [])
    pl.setdefault('notes', [])

    def issue(level, text):
        if not any(q['text'] == text for q in T['qualityIssues']):
            T['qualityIssues'].append(dict(level=level, text=text))
            issues.append((level, text))

    def pl_note(level, text):
        if not any(n['text'] == text for n in pl['notes']):
            pl['notes'].append(dict(level=level, text=text))
            return True
        return False

    RS = Resolver(data, roster)
    cats_norm = {norm(c): c for c in T['categorias']}

    # --- roster: nombre / equipo de los jugadores que ya existen
    # Un nombre repetido en el roster bajo ≥2 IDs no se usa para renombrar: son jugadores distintos.
    if roster:
        for pid, r in roster.items():
            p = P.get(pid)
            if not p or p.get('isUnresolvedId'):
                continue
            if RS.is_dup_name(r['name']) and norm(p['name']) != norm(r['name']):
                continue                                   # el roster le pone el nombre de otra persona: no se toca
            if norm(p['name']) != norm(r['name']):
                changes.append(f'Roster: {pid} cambia de nombre «{p["name"]}» → «{r["name"]}»')
                p['name'] = r['name']
            if r['team'] and p.get('equipoOficial') != r['team']:
                changes.append(f'Roster: {pid} ({p["name"]}) equipo «{p.get("equipoOficial")}» → «{r["team"]}»')
                p['equipoOficial'] = r['team']
            p['categoriaConflict'] = bool(p.get('equipoOficial') and p['equipoOficial'] != p['anthro']['categoria'])

    # --- pliegues: asociar nombre -> ID (antes de las alturas, porque algunos archivos traen altura)
    pl_items = []           # (fila, id, cómo)
    pl_skipped = []         # (nivel, texto)
    for w in pliegues or []:
        hint = w['categoria'] or w['hint']
        if w['id']:
            rid, how = w['id'], 'id'
        else:
            rid, how, cand = RS.lookup(w['name'], hint)
            if rid is None:
                where = f'«{w["name"]}» ({w["sheet"]}, {fmt_date(w["fecha"])})'
                if how == 'ambiguo':
                    pl_skipped.append(('warn', f'Pliegues {where}: el nombre coincide con varios IDs ({", ".join(cand)}) — no se ha vinculado. Pon su ID RCF_xxx en el Excel de pliegues.'))
                else:
                    pl_skipped.append(('warn', f'Pliegues {where}: no se encuentra en el roster — no se ha integrado. Pon su ID RCF_xxx (o corrige el nombre) en el Excel de pliegues.'))
                continue
        pl_items.append((w, rid, how))

    registro_list = list(registro_list)
    team_of = lambda pid: (roster or {}).get(pid, {}).get('team')
    for w, rid, how in pl_items:                      # alturas que vienen en el archivo de pliegues
        if w['altura'] is None:
            continue
        cat = cats_norm.get(norm(w['categoria'] or '')) or (team_of(rid) if team_of(rid) in T['categorias'] else None)
        registro_list.append(dict(sheet=w['sheet'], id=rid, fecha=w['fecha'], categoria=cat if rid not in P else None,
                                  altura=w['altura'], peso=None, vals={}, fuente='Pliegues/altura ' + fmt_date(w['fecha'])))

    # --- mediciones nuevas
    cat_latest = {}
    for q in data['players']:
        lk = latest_known(q)
        if lk and (q['anthro']['categoria'] not in cat_latest or lk > cat_latest[q['anthro']['categoria']]):
            cat_latest[q['anthro']['categoria']] = lk
    by_id = {}
    for r in registro_list:
        r['categoria'] = cats_norm.get(norm(r['categoria'] or ''), r['categoria'])
        rid = RS.resolve_id(r['id'], r['categoria'])[0]
        if rid is None:
            if r['fecha'] <= cat_latest.get(r['categoria'], dt.date.min):
                continue                      # fila antigua sin ID: ya estaba integrada en su día
            issue('warn', f'Fila «{r["id"]}» ({r["sheet"]}, {fmt_date(r["fecha"])}): sin ID RCF y no se pudo vincular con ningún jugador por nombre — no se ha integrado. Pon su ID RCF_xxx en el Excel.')
            continue
        r['id'] = rid
        by_id.setdefault(rid, []).append(r)
    new_dates = set()
    created_from_pl = {}
    for pid, rows in by_id.items():
        rows.sort(key=lambda r: r['fecha'])
        for r in rows:
            p = P.get(pid)
            if p is not None:
                lk = latest_known(p)
                if lk and r['fecha'] <= lk:
                    continue
            has_perim = bool(r['vals'])
            if r['altura'] is None and not has_perim and r['peso'] is None:
                continue
            ds = fmt_date(r['fecha'])
            # jugador nuevo
            if p is None:
                name = roster.get(pid, {}).get('name') if roster else None
                team = roster.get(pid, {}).get('team') if roster else None
                cat = r['categoria'] or (team if team in T['categorias'] else None)
                if cat is None:
                    issue('warn', f'{pid} aparece en «{r["sheet"]}» ({ds}) pero no está en el roster ni trae categoría — no se ha podido crear su ficha.')
                    continue
                tmpl_src = [q for q in data['players'] if q['anthro']['categoria'] == cat and not q.get('hasVald')] or \
                           [q for q in data['players'] if not q.get('hasVald')]
                p = copy.deepcopy(tmpl_src[0])
                stage = Counter(q['phvStage'] for q in data['players'] if q['anthro']['categoria'] == cat).most_common(1)
                p.update(id=pid, name=name or f'{pid} (sin nombre en el roster)', equipoOficial=team or cat,
                         categoriaConflict=bool(team and team != cat), isUnresolvedId=False, pliegues=None, hasPliegues=False,
                         pctJump=None, pctRsi=None, pctLandingSym=None, pctConcSym=None, hasVald=False,
                         growth=None, growthDiscarded=None, phvStage=stage[0][0] if stage else None, phvAlert=None,
                         riskFlags=[], riskCount=0, heightOutlier=None, alturas=[])
                for k in [k for k in p if k.startswith('vald')]:
                    del p[k]
                p['anthro'] = dict(fecha=ds, categoria=cat, altura=None, peso=None, pesoMissing=True,
                                   raw={k: None for k in RAW_KEYS}, **{f: None for f in ('pct_altura', 'pct_peso', 'pct_muslo_medio_avg', 'pct_dorsiflex_avg', 'pct_cadera', 'pct_pecho')})
                derive_anthro(p['anthro'])
                data['players'].append(p)
                P[pid] = p
                if not name:
                    issue('warn', f'{pid} aparece en el Registro ({r["sheet"]}) pero no está en el roster ID_JUGADOR_RCF: se muestra con identificador provisional.')
                if str(r.get('fuente', '')).startswith('Pliegues'):
                    created_from_pl[pid] = ds
                    txt = f'{pid} ({p["name"]}) no estaba en el test de antropometría — se añade a {cat} con la altura y los pliegues del {ds}; el resto de medidas siguen pendientes.'
                    issue('info', txt)
                    pl_note('info', txt)
                changes.append(f'Jugador nuevo: {pid} {p["name"]} ({cat})')
                old_h, old_d, first = None, None, True
            else:
                first = False
                a0 = p['anthro']
                old_h, old_d = a0['altura'], a0.get('alturaFecha') or a0['fecha']
                if r['categoria'] and r['categoria'] != a0['categoria']:
                    issue('warn', f'{pid} ({p["name"]}): la hoja {r["sheet"]} lo registra como «{r["categoria"]}» pero el dashboard lo tiene en «{a0["categoria"]}» — se mantiene la categoría del dashboard.')
            a = p['anthro']
            if has_perim:                                     # medición completa
                if not first:
                    a['alturaPrev'] = dict(valor=old_h, fecha=old_d) if old_h is not None else a.get('alturaPrev')
                a['fecha'] = ds
                a.pop('alturaFecha', None)
                a['raw'] = {k: r['vals'].get(k) for k in RAW_KEYS}
                if r['altura'] is not None:
                    a['altura'] = r['altura']
                derive_anthro(a)
                if r['peso'] is not None:
                    a['peso'], a['pesoMissing'] = r['peso'], False
                p['riskFlags'] = antro_flags(a) + [f for f in p['riskFlags'] if f.get('source') != 'antro']
                p['riskCount'] = len(p['riskFlags'])
                kind = 'medición completa'
            else:                                              # solo altura (y/o peso)
                if r['altura'] is not None:
                    if not first and old_h is not None:
                        a['alturaPrev'] = dict(valor=old_h, fecha=old_d)
                    a['altura'] = r['altura']
                    if first:
                        a['fecha'] = ds
                    else:
                        a['alturaFecha'] = ds
                if r['peso'] is not None:
                    a['peso'], a['pesoMissing'] = r['peso'], False
                kind = 'altura'
            # histórico de alturas
            if r['altura'] is not None:
                hist = {x['fecha']: x for x in p.get('alturas') or []}
                g0 = p.get('growth')
                if g0:
                    hist.setdefault(g0['prevDate'], dict(fecha=g0['prevDate'], valor=g0['prevHeight'], fuente=f'Histórico {T.get("historicoFecha", "")}'))
                if old_h is not None:
                    hist.setdefault(old_d, dict(fecha=old_d, valor=old_h, fuente='Test antropometría'))
                hist[ds] = dict(fecha=ds, valor=r['altura'], fuente=r.get('fuente') or 'Registro Valoraciones ' + ds)
                p['alturas'] = sorted(hist.values(), key=lambda x: dmy(x['fecha']))
                # avisos de calidad
                if old_h is not None:
                    dlt = round(r['altura'] - old_h, 1)
                    days = (r['fecha'] - dmy(old_d)).days
                    if dlt < -1.0:
                        issue('warn', f'Altura {pid} ({p["name"]}): {old_h} cm ({old_d}) → {r["altura"]} cm ({ds}) = {dlt} cm. Un descenso no es fisiológico: revisar si es error de medición/registro.')
                    elif dlt > 3.0 and days < 100:
                        issue('warn', f'Altura {pid} ({p["name"]}): {old_h} cm ({old_d}) → {r["altura"]} cm ({ds}) = +{dlt} cm en {days} días — muy alto, confirmar la medida.')
            recompute_growth(p, thr)
            touched_cats.add(a['categoria'])
            new_dates.add(ds)
            changes.append(f'{pid} {p["name"]} · {a["categoria"]} · {ds} · {kind}' + (f' · altura {r["altura"]}' if r['altura'] is not None else ''))

    # --- pliegues cutáneos
    standalone = {}                                  # id -> registro sin ficha individual
    for rec in pl['records']:
        if not rec.get('playerId'):
            o = standalone.get(rec['id'])
            if o is None or dmy(rec['fecha']) > dmy(o['fecha']):
                standalone[rec['id']] = rec
    for pid in [i for i in standalone if i in P]:    # ya tiene ficha: los pliegues pasan a la ficha
        rec, p = standalone.pop(pid), P[pid]
        if not p.get('pliegues') or dmy(rec['fecha']) > dmy(p['pliegues']['fecha']):
            p['pliegues'] = dict(fecha=rec['fecha'], **{k: rec[k] for k in FOLDS}, sumatorio=rec['sum8'],
                                 sumatorioArchivo=rec['sum8'], pctFatEst=rec['pctFat'], matchedBy=rec['matchedBy'])
            p['hasPliegues'] = True
            changes.append(f'Pliegues {pid} {p["name"]}: pasan a su ficha individual')
    pl_by_id = {}
    for w, rid, how in pl_items:
        pl_by_id.setdefault(rid, []).append((w, how))
    linked = {}                                       # (fecha, hint) -> [(nombre en archivo, id, aproximado)]
    for rid, lst in sorted(pl_by_id.items()):
        lst.sort(key=lambda x: x[0]['fecha'])
        for w, how in lst:
            p = P.get(rid)
            cur = None
            if p is not None and p.get('pliegues'):
                cur = dmy(p['pliegues']['fecha'])
            elif rid in standalone:
                cur = dmy(standalone[rid]['fecha'])
            if cur and w['fecha'] <= cur:
                continue                              # ya integrado
            ds = fmt_date(w['fecha'])
            who = (p['name'] if p else RS.display(rid)) or w['name'] or rid
            if any(w['folds'][k] is None for k in FOLDS):
                pl_skipped.append(('warn', f'Pliegues {who} ({rid}, {ds}): faltan pliegues en el archivo — no se ha integrado hasta que esté completo.'))
                continue
            f = {k: w['folds'][k] for k in FOLDS}
            s8 = round(sum(f.values()), 1)
            fat = round(0.735 * (f['triceps'] + f['piernaMedial']) + 1.0, 1)
            if w['sumFile'] is not None and abs(s8 - w['sumFile']) > 0.05:
                if abs(s8 - f['triceps'] - w['sumFile']) <= 0.05:
                    if not any('NO incluye el tríceps' in n['text'] for n in pl['notes']):
                        pl_note('warn', f'El archivo de pliegues de {w["hint"] or w["sheet"]} trae una columna "Sumatorio" que NO incluye el tríceps (suma 7 pliegues). Aquí se usa Σ8 completo, así que difiere de lo que muestra ese Excel.')
                else:
                    txt = f'Pliegues {ds}: {who} — el sumatorio del archivo ({w["sumFile"]}) no coincide con la suma de los 8 pliegues ({s8}); el dashboard usa {s8}.'
                    pl_note('warn', txt)
                    issue('warn', txt)
            mb = 'id' if how == 'id' else how
            if p is not None:
                p['pliegues'] = dict(fecha=ds, **f, sumatorio=s8, sumatorioArchivo=w['sumFile'] if w['sumFile'] is not None else s8,
                                     pctFatEst=fat, matchedBy=mb)
                p['hasPliegues'] = True
                standalone.pop(rid, None)
            else:
                cat = cats_norm.get(norm(w['categoria'] or '')) or (team_of(rid) if team_of(rid) in T['categorias'] else None)
                if cat is None:
                    cat = next((c for c in T['categorias'] if w['hint'] and norm(c).startswith(norm(w['hint']))), None) or w['categoria'] or w['hint'] or 'Sin categoría'
                standalone[rid] = dict(playerId=None, id=rid, name=RS.display(rid) or w['name'] or rid, categoria=cat, fecha=ds,
                                       altura=None, alturaFecha=None, alturaPrev=None, **f, sum8=s8, pctFat=fat, matchedBy=mb)
            if how != 'id':
                linked.setdefault((ds, w['hint']), []).append((w['name'], rid, how == 'name'))
            changes.append(f'Pliegues {rid} {who} · {ds}' + (f' (vinculado por nombre «{w["name"]}»)' if how != 'id' else ''))
    for (ds, hint), lst in linked.items():
        approx = any(x[2] for x in lst)
        txt = (f'Pliegues {ds}{f" ({hint})" if hint else ""}: {len(lst)} {"jugador identificado" if len(lst) == 1 else "jugadores identificados"} solo por nombre en el Excel → vinculado{"" if len(lst) == 1 else "s"} automáticamente con su ID del roster: ' +
               '; '.join(f'{n} → {i}{" (aprox.)" if ap else ""}' for n, i, ap in lst) + '.')
        pl_note('warn' if approx else 'info', txt)

    # --- avisos regenerados en cada sincronización (desaparecen solos cuando se corrigen en el Excel)
    dup_items = []
    for key, ids in RS.dup_groups.items():
        disp = RS.display(ids[0])
        extra = [f'{i} figura como «{P[i]["name"]}» en el dashboard' for i in ids if i in P and norm(P[i]['name']) != norm(disp)]
        cats = sorted({c for i in ids for c in ((P[i]['anthro']['categoria'],) if i in P else ((roster or {}).get(i, {}).get('team'),)) if c})
        dup_items.append(('warn', f'Roster: el nombre «{disp}» está repetido en {" y ".join(ids)}' + (f' ({", ".join(cats)})' if cats else '') +
                          '. Se tratan como personas distintas: no se renombran ni se vinculan por nombre' + (f' ({"; ".join(extra)})' if extra else '') +
                          '. Para vincularlos desde un Excel por nombre, pon su ID RCF_xxx o diferencia el nombre en el roster.'))
    T['qualityIssues'] = [q for q in T['qualityIssues'] if not q['text'].startswith('Roster: 3 pares de IDs con nombre idéntico')]
    if _set_auto(T['qualityIssues'], 'dupname', dup_items):
        issues.append(('warn', 'avisos de nombres repetidos actualizados'))
    skipped_changed = _set_auto(T['qualityIssues'], 'pl-skipped', pl_skipped)
    _set_auto(pl['notes'], 'pl-skipped', pl_skipped)
    if skipped_changed:
        issues.append(('warn', 'avisos de pliegues no integrados actualizados'))
    pl['notes'] = [n for n in pl['notes'] if not n['text'].startswith('Registros RCF_074, RCF_087, RCF_093')]
    sa_items = []
    if standalone:
        sa_items.append(('warn', 'Registros con pliegues pero sin antropometría en el dashboard (aparecen en esta pestaña sin ficha individual): ' +
                         '; '.join(f'{r["name"]} ({r["id"]}, {r["categoria"]}, {r["fecha"]})' for r in sorted(standalone.values(), key=lambda r: (r['categoria'], r['name']))) + '.'))
    _set_auto(pl['notes'], 'pl-standalone', sa_items)

    if not changes and not issues:
        return []
    if not changes:
        changes = ['(solo avisos de calidad nuevos)']

    # --- percentiles, outliers y curva de crecimiento de las categorías tocadas
    for cat in touched_cats:
        members = [p for p in data['players'] if p['anthro']['categoria'] == cat]
        for fld, src in [('pct_altura', 'altura'), ('pct_muslo_medio_avg', 'muslo_medio_avg'), ('pct_dorsiflex_avg', 'dorsiflex_avg'),
                         ('pct_cadera', 'cadera'), ('pct_pecho', 'pecho'), ('pct_peso', 'peso')]:
            vals = [q['anthro'][src] for q in members if q['anthro'].get(src) is not None]
            for q in members:
                v = q['anthro'].get(src)
                q['anthro'][fld] = pct_rank(vals, v) if v is not None else None
        hs = [q['anthro']['altura'] for q in members if q['anthro']['altura'] is not None]
        for q in members:
            pc = q['anthro']['pct_altura']
            q['heightOutlier'] = None if pc is None else ('alto' if pc >= 90 else 'bajo' if pc <= 10 else None)
        if hs:
            arr = np.array(hs)
            g = next((x for x in T['growthCurve'] if x['categoria'] == cat), None)
            if g is None:
                g = dict(categoria=cat)
                T['growthCurve'].append(g)
            g.update(n=len(arr), mean=round(float(arr.mean()), 1), median=round(float(np.median(arr)), 1),
                     min=round(float(arr.min()), 1), max=round(float(arr.max()), 1),
                     p10=round(float(np.percentile(arr, 10)), 1), p25=round(float(np.percentile(arr, 25)), 1),
                     p75=round(float(np.percentile(arr, 75)), 1), p90=round(float(np.percentile(arr, 90)), 1))

    # --- registros de la pestaña Pliegues (se reconstruyen desde las fichas + los registros sin ficha)
    cats = T['categorias']
    recs = [_pl_record_from_player(p) for p in data['players'] if p.get('hasPliegues') and p.get('pliegues')] + list(standalone.values())
    recs.sort(key=lambda r: (cats.index(r['categoria']) if r['categoria'] in cats else 99, r['name']))
    pl['records'] = recs
    T['unmatchedPliegues'] = []

    # --- equipo
    T['nPlayers'] = len(data['players'])
    hs_all = [p['anthro']['altura'] for p in data['players'] if p['anthro']['altura'] is not None]
    T['avgHeight'] = round(sum(hs_all) / len(hs_all), 1)
    T['nPhvAlerts'] = sum(1 for p in data['players'] if p['phvAlert'])
    T['fechasAntro'] = sorted(set(T['fechasAntro']) | new_dates, key=dmy)
    T['generatedAt'] = now_iso()
    if 'Sincronización automática' not in T.get('dataSource', ''):
        T['dataSource'] = T.get('dataSource', '') + ' · Sincronización automática con Drive'
    return changes


# ----------------------------------------------------------------------------- Drive
def drive_download_xlsx(file_id, sa_info):
    from google.oauth2 import service_account
    from google.auth.transport.requests import AuthorizedSession
    creds = service_account.Credentials.from_service_account_info(
        sa_info, scopes=['https://www.googleapis.com/auth/drive.readonly'])
    s = AuthorizedSession(creds)
    meta = s.get(f'https://www.googleapis.com/drive/v3/files/{file_id}', params={'fields': 'mimeType,name', 'supportsAllDrives': 'true'}, timeout=60)
    meta.raise_for_status()
    mt = meta.json()['mimeType']
    if mt == 'application/vnd.google-apps.spreadsheet':
        r = s.get(f'https://www.googleapis.com/drive/v3/files/{file_id}/export',
                  params={'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}, timeout=120)
    else:
        r = s.get(f'https://www.googleapis.com/drive/v3/files/{file_id}', params={'alt': 'media', 'supportsAllDrives': 'true'}, timeout=120)
    r.raise_for_status()
    return r.content


# ----------------------------------------------------------------------------- dataset <-> html
def load_dataset(html):
    m = DATASET_RE.search(html)
    if not m:
        raise SystemExit('No encuentro el bloque <script id="dataset"> en el HTML')
    return json.loads(m.group(2)), m


def dump_html(html, m, data):
    ds = json.dumps(clean(data), ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    return html[:m.start(2)] + ds + html[m.end(2):]


def run(html_text, registros, roster, log=print, stamp=False, pliegues=None):
    """registros: lista de bytes/paths de Excel; roster: bytes/path o None;
    pliegues: lista de (bytes/path, categoría_pista) de los Excel de pliegues.
    Devuelve (html_nuevo|None, cambios). Con stamp=True, aunque no haya cambios devuelve el HTML con la
    hora de sincronización actualizada (para la app en vivo)."""
    data, m = load_dataset(html_text)
    rows = []
    for rg in registros:
        rows += read_registro(io.BytesIO(rg) if isinstance(rg, (bytes, bytearray)) else rg, log)
    ros = read_roster(io.BytesIO(roster) if isinstance(roster, (bytes, bytearray)) else roster) if roster else None
    pl_rows = []
    for src, hint in pliegues or []:
        pl_rows += read_pliegues(io.BytesIO(src) if isinstance(src, (bytes, bytearray)) else src, hint, log)
    changes = sync(data, rows, ros, log, pliegues=pl_rows)
    if not changes and not stamp:
        return None, []
    if stamp:
        data['team']['generatedAt'] = dt.datetime.now().strftime('%d/%m/%Y %H:%M')
    return dump_html(html_text, m, data), changes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--html', default='dashboard.html')
    ap.add_argument('--out')
    ap.add_argument('--registro', action='append', default=[])
    ap.add_argument('--roster')
    ap.add_argument('--pliegues', action='append', default=[], help='ruta[=categoría]  (p. ej. infantil.xlsx=Infantil)')
    ap.add_argument('--config')
    ap.add_argument('--sa-json', help='ruta al JSON de la cuenta de servicio (o variable GOOGLE_SERVICE_ACCOUNT_JSON)')
    a = ap.parse_args()
    html = open(a.html, encoding='utf-8').read()
    regs, roster = list(a.registro), a.roster
    pls = [tuple((x.rsplit('=', 1) + [None])[:2]) if '=' in x else (x, None) for x in a.pliegues]
    if a.config:
        cfg = json.load(open(a.config))
        sa = a.sa_json and open(a.sa_json).read() or os.environ.get('GOOGLE_SERVICE_ACCOUNT_JSON')
        if not sa:
            raise SystemExit('Falta la cuenta de servicio (--sa-json o GOOGLE_SERVICE_ACCOUNT_JSON)')
        sa = json.loads(sa)
        regs = [drive_download_xlsx(i, sa) for i in cfg['registro_ids']]
        roster = drive_download_xlsx(cfg['roster_id'], sa) if cfg.get('roster_id') else None
        pls = [(drive_download_xlsx(x['id'], sa), x.get('categoria')) for x in cfg.get('pliegues', [])]
    new_html, changes = run(html, regs, roster, pliegues=pls)
    if not new_html:
        print('Sin cambios: el dashboard ya está al día.')
        return 0
    open(a.out or a.html, 'w', encoding='utf-8').write(new_html)
    print(f'{len(changes)} cambios:')
    for c in changes:
        print(' -', c)
    return 0


if __name__ == '__main__':
    sys.exit(main())
